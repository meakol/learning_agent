from __future__ import annotations

import datetime as dt
import hashlib
import math
import re
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

try:  # PyMuPDF 1.24+ exposes both names
    import pymupdf  # type: ignore
except ImportError:  # pragma: no cover
    import fitz as pymupdf  # type: ignore

DESIGN_WIDTH = 1280.0
MAX_TEXT_PER_SLIDE = 4000


class ParseError(Exception):
    pass


def _now() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def _clean_line(text: str) -> str:
    text = text.replace("\u000b", " ").replace("\x0b", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _hex_color(color_format: Any) -> str | None:
    try:
        if color_format is None or color_format.type is None:
            return None
        rgb = color_format.rgb
        if rgb is None:
            return None
        return f"#{str(rgb).lower()}"
    except Exception:
        return None


def _is_dark(hex_color: str | None) -> bool:
    if not hex_color or not re.fullmatch(r"#[0-9a-f]{6}", hex_color):
        return False
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (1, 3, 5))
    return (0.299 * r + 0.587 * g + 0.114 * b) < 110


def _luminance(hex_color: str | None) -> int | None:
    if not hex_color or not re.fullmatch(r"#[0-9a-f]{6}", hex_color):
        return None
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (1, 3, 5))
    return round(0.299 * r + 0.587 * g + 0.114 * b)


def _solid_fill(shape: Any) -> str | None:
    try:
        fill = shape.fill
        if fill.type == MSO_FILL.SOLID:
            return _hex_color(fill.fore_color)
    except Exception:
        pass
    return None


def _line_color(shape: Any) -> str | None:
    try:
        line = shape.line
        if line.width is not None:
            return _hex_color(line.color)
    except Exception:
        pass
    return None


def _shape_bg(slide: Any) -> str | None:
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            fill = owner.background.fill
            if fill.type == MSO_FILL.SOLID:
                color = _hex_color(fill.fore_color)
                if color:
                    return color
        except Exception:
            continue
    return None


def _walk_shapes(
    shapes: Iterable[Any],
    origin_x: float = 0.0,
    origin_y: float = 0.0,
    scale_x: float = 1.0,
    scale_y: float = 1.0,
) -> Iterable[tuple[Any, float, float, float, float]]:
    """Yield shapes with absolute EMU coordinates, flattening groups."""
    for shape in shapes:
        try:
            left = shape.left or 0
            top = shape.top or 0
            width = shape.width or 0
            height = shape.height or 0
        except Exception:
            continue
        x = origin_x + left * scale_x
        y = origin_y + top * scale_y
        w = width * scale_x
        h = height * scale_y
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                xfrm = shape._element.grpSpPr.xfrm
                ch_off, ch_ext = xfrm.chOff, xfrm.chExt
                child_sx = width / ch_ext.cx if ch_ext.cx else 1.0
                child_sy = height / ch_ext.cy if ch_ext.cy else 1.0
                yield from _walk_shapes(
                    shape.shapes,
                    x - ch_off.x * scale_x * child_sx,
                    y - ch_off.y * scale_y * child_sy,
                    scale_x * child_sx,
                    scale_y * child_sy,
                )
                continue
            except Exception:
                pass
        yield shape, x, y, w, h


def _paragraph_bullet(paragraph: Any, is_body: bool) -> str | None:
    pPr = paragraph._p.find(qn("a:pPr"))
    if pPr is not None:
        if pPr.find(qn("a:buNone")) is not None:
            return None
        bu_char = pPr.find(qn("a:buChar"))
        if bu_char is not None:
            return bu_char.get("char") or "•"
        if pPr.find(qn("a:buAutoNum")) is not None:
            return "#"
    if is_body:
        return "•"
    return None


def _alignment(paragraph: Any) -> str:
    align = paragraph.alignment
    if align == PP_ALIGN.CENTER:
        return "center"
    if align == PP_ALIGN.RIGHT:
        return "right"
    if align == PP_ALIGN.JUSTIFY:
        return "justify"
    return "left"


def _vertical_align(text_frame: Any) -> str:
    anchor = text_frame.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        return "center"
    if anchor == MSO_ANCHOR.BOTTOM:
        return "flex-end"
    return "flex-start"


def _font_color(run: Any) -> str | None:
    try:
        return _hex_color(run.font.color)
    except Exception:
        return None


def _text_element(
    shape: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    px_per_emu: float,
    font_px_per_pt: float,
    is_title: bool,
    default_color: str,
) -> dict[str, Any] | None:
    text_frame = shape.text_frame
    try:
        is_body = shape.is_placeholder and shape.placeholder_format.type in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
        )
    except Exception:
        is_body = False
    try:
        is_subtitle = shape.is_placeholder and shape.placeholder_format.type in (
            PP_PLACEHOLDER.SUBTITLE,
        )
    except Exception:
        is_subtitle = False

    paragraphs: list[dict[str, Any]] = []
    for paragraph in text_frame.paragraphs:
        runs: list[dict[str, Any]] = []
        for run in paragraph.runs:
            text = run.text
            if not text:
                continue
            size_pt = run.font.size.pt if run.font.size else None
            if size_pt is None:
                size_pt = 30 if is_title else 20 if is_subtitle else 18
            runs.append(
                {
                    "text": text,
                    "bold": bool(run.font.bold),
                    "italic": bool(run.font.italic),
                    "underline": bool(run.font.underline),
                    "size": round(size_pt * font_px_per_pt, 1),
                    "color": _font_color(run),
                    "font": run.font.name,
                }
            )
        if not runs:
            text = _clean_line(paragraph.text)
            if text:
                size_pt = 30 if is_title else 20 if is_subtitle else 18
                runs.append(
                    {
                        "text": text,
                        "bold": is_title,
                        "italic": False,
                        "underline": False,
                        "size": round(size_pt * font_px_per_pt, 1),
                        "color": None,
                        "font": None,
                    }
                )
        if not runs:
            continue
        paragraphs.append(
            {
                "runs": runs,
                "align": _alignment(paragraph),
                "level": min(int(paragraph.level or 0), 5),
                "bullet": _paragraph_bullet(paragraph, is_body),
                "space_after": round(
                    (paragraph.space_after.pt if paragraph.space_after else 0)
                    * font_px_per_pt,
                    1,
                ),
            }
        )
    if not paragraphs:
        return None
    return {
        "type": "text",
        "x": round(x * px_per_emu, 1),
        "y": round(y * px_per_emu, 1),
        "w": round(w * px_per_emu, 1),
        "h": round(h * px_per_emu, 1),
        "fill": _solid_fill(shape),
        "line": _line_color(shape),
        "valign": _vertical_align(text_frame),
        "paragraphs": paragraphs,
        "default_color": default_color,
    }


def _shape_label(shape: Any) -> str | None:
    if getattr(shape, "has_chart", False):
        try:
            title = shape.chart.chart_title.text_frame.text.strip()
        except Exception:
            title = ""
        return f"图表{('：' + title) if title else ''}"
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
            return "嵌入对象"
    except Exception:
        pass
    return None


def _collect_text(shapes: Iterable[Any]) -> str:
    lines: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                lines.append(_collect_text(shape.shapes))
                continue
            except Exception:
                pass
        try:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = _clean_line(paragraph.text)
                    if text:
                        lines.append(text)
        except Exception:
            pass
        try:
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [_clean_line(cell.text) for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        except Exception:
            pass
        try:
            if getattr(shape, "has_chart", False):
                try:
                    chart_title = shape.chart.chart_title.text_frame.text.strip()
                except Exception:
                    chart_title = ""
                if chart_title:
                    lines.append(chart_title)
        except Exception:
            pass
    return "\n".join(line for line in lines if line)


def _notes_text(slide: Any) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        return _clean_line(slide.notes_slide.notes_text_frame.text)
    except Exception:
        return ""


def _slide_title(slide: Any) -> str:
    try:
        title_shape = slide.shapes.title
        if title_shape is not None:
            title = re.sub(r"\s*\n\s*", " ", _clean_line(title_shape.text))
            if title:
                return title[:100]
    except Exception:
        pass

    # No title placeholder (common in designed decks): use the largest text
    # on the slide, which is far more reliable than shape order.
    best_text = ""
    best_size = -1.0
    for shape in slide.shapes:
        try:
            if not shape.has_text_frame:
                continue
            sizes: list[float] = []
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size:
                        sizes.append(run.font.size.pt)
            size = max(sizes) if sizes else 0.0
            text = _clean_line(
                " ".join(
                    line for line in shape.text_frame.text.splitlines() if line.strip()
                )
            )
            if not text:
                continue
            if size > best_size:
                best_size = size
                best_text = text
        except Exception:
            continue
    if best_text:
        return best_text[:100]

    for shape in slide.shapes:
        try:
            if shape.has_text_frame:
                first_line = _clean_line(shape.text_frame.text.splitlines()[0])
                if first_line:
                    return first_line[:80]
        except Exception:
            continue
    return ""


# --------------------------------------------------------------------------- #
# pptx
# --------------------------------------------------------------------------- #
def _parse_pptx(path: Path, doc_id: str, media_dir: Path) -> dict[str, Any]:
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法读取 PPTX：{exc}") from exc

    slides: list[dict[str, Any]] = []
    warnings: list[str] = []
    slide_width = int(prs.slide_width or 12192000)
    slide_height = int(prs.slide_height or 6858000)
    px_per_emu = DESIGN_WIDTH / slide_width
    font_px_per_pt = DESIGN_WIDTH / (slide_width / 914400 * 72)
    image_hash_counts: dict[str, int] = {}

    for index, slide in enumerate(prs.slides, start=1):
        if index > 200:
            warnings.append("课件页数较多，仅解析前 200 页。")
            break

        bg = _shape_bg(slide)
        elements: list[dict[str, Any]] = []
        images = 0
        color_samples: list[str] = []
        shape_index = 0

        for shape, x, y, w, h in _walk_shapes(slide.shapes):
            shape_index += 1
            if w <= 1 or h <= 1:
                continue
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        ext = (image.ext or "png").lower()
                        if ext in ("emf", "wmf", "svg"):
                            elements.append(
                                {
                                    "type": "placeholder",
                                    "x": round(x * px_per_emu, 1),
                                    "y": round(y * px_per_emu, 1),
                                    "w": round(w * px_per_emu, 1),
                                    "h": round(h * px_per_emu, 1),
                                    "label": "矢量图片",
                                }
                            )
                            continue
                        name = f"s{index}_{shape_index}.{ext}"
                        blob = image.blob
                        (media_dir / name).write_bytes(blob)
                        digest = hashlib.sha1(blob).hexdigest()
                        image_hash_counts[digest] = image_hash_counts.get(digest, 0) + 1
                        elements.append(
                            {
                                "type": "image",
                                "x": round(x * px_per_emu, 1),
                                "y": round(y * px_per_emu, 1),
                                "w": round(w * px_per_emu, 1),
                                "h": round(h * px_per_emu, 1),
                                "url": f"/api/doc/{doc_id}/media/{name}",
                                "_media_hash": digest,
                            }
                        )
                        images += 1
                        continue
                    except Exception:
                        pass

                if shape.has_table:
                    table = shape.table
                    rows = [
                        [_clean_line(cell.text) for cell in row.cells]
                        for row in table.rows
                    ]
                    elements.append(
                        {
                            "type": "table",
                            "x": round(x * px_per_emu, 1),
                            "y": round(y * px_per_emu, 1),
                            "w": round(w * px_per_emu, 1),
                            "h": round(h * px_per_emu, 1),
                            "rows": rows,
                            "col_widths": [
                                round(col.width * px_per_emu, 1)
                                for col in table.columns
                            ],
                            "row_heights": [
                                round(row.height * px_per_emu, 1)
                                for row in table.rows
                            ],
                            "header": bool(getattr(table, "first_row", False)),
                        }
                    )
                    continue

                is_title = False
                try:
                    is_title = slide.shapes.title is not None and shape == slide.shapes.title
                except Exception:
                    pass

                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            color = _font_color(run)
                            if color:
                                color_samples.append(color)
                    element = _text_element(
                        shape,
                        x,
                        y,
                        w,
                        h,
                        px_per_emu,
                        font_px_per_pt,
                        is_title,
                        "#1f2937",
                    )
                    if element:
                        try:
                            shape_name = str(shape.auto_shape_type)
                        except Exception:
                            shape_name = ""
                        element["radius"] = (
                            "50%"
                            if "OVAL" in shape_name
                            else "10px"
                            if "ROUND" in shape_name
                            else "0"
                        )
                        try:
                            element["rotation"] = float(shape.rotation or 0)
                        except Exception:
                            element["rotation"] = 0
                        elements.append(element)
                        continue

                label = _shape_label(shape)
                if label:
                    elements.append(
                        {
                            "type": "placeholder",
                            "x": round(x * px_per_emu, 1),
                            "y": round(y * px_per_emu, 1),
                            "w": round(w * px_per_emu, 1),
                            "h": round(h * px_per_emu, 1),
                            "label": label,
                        }
                    )
                    continue

                fill = _solid_fill(shape)
                line = _line_color(shape)
                if fill or line:
                    try:
                        shape_name = str(shape.auto_shape_type)
                    except Exception:
                        shape_name = ""
                    elements.append(
                        {
                            "type": "shape",
                            "x": round(x * px_per_emu, 1),
                            "y": round(y * px_per_emu, 1),
                            "w": round(w * px_per_emu, 1),
                            "h": round(h * px_per_emu, 1),
                            "fill": fill,
                            "line": line,
                            "radius": (
                                "50%"
                                if "OVAL" in shape_name
                                else "10px"
                                if "ROUND" in shape_name
                                else "0"
                            ),
                        }
                    )
            except Exception:
                continue

        # Adapt text color when the slide is a dark deck.
        light_colors = 0
        dark_colors = 0
        for color in color_samples:
            lum = _luminance(color)
            if lum is None:
                continue
            if lum > 190:
                light_colors += 1
            elif lum < 90:
                dark_colors += 1
        if bg is None and light_colors >= 3 and light_colors > dark_colors * 1.5:
            bg = "#111827"

        default_color = "#e5e7eb" if _is_dark(bg) else "#1f2937"
        for element in elements:
            if element.get("type") == "text":
                element["default_color"] = default_color
                for paragraph in element["paragraphs"]:
                    for run in paragraph["runs"]:
                        if not run.get("color"):
                            run["color"] = default_color

        title = _slide_title(slide)
        text = _collect_text(slide.shapes)
        notes = _notes_text(slide)
        if not text and images == 0 and not any(
            e["type"] in ("table", "placeholder") for e in elements
        ):
            continue
        if len(text) > MAX_TEXT_PER_SLIDE:
            text = text[:MAX_TEXT_PER_SLIDE] + "…"

        slide_height_px = DESIGN_WIDTH * slide_height / slide_width
        slide_area = DESIGN_WIDTH * slide_height_px
        text_elements = [element for element in elements if element.get("type") == "text"]

        def nearby_text_count(x: float, y: float, w: float, h: float, margin: float = 110.0) -> int:
            count = 0
            for text_element in text_elements:
                ex = float(text_element.get("x") or 0)
                ey = float(text_element.get("y") or 0)
                ew = float(text_element.get("w") or 0)
                eh = float(text_element.get("h") or 0)
                if ew < 12 or eh < 8:
                    continue
                # Ignore title / header / footer when judging figure context.
                if ey < slide_height_px * 0.11 or ey > slide_height_px * 0.92:
                    continue
                if (
                    ex < x + w + margin
                    and ex + ew > x - margin
                    and ey < y + h + margin
                    and ey + eh > y - margin
                ):
                    count += 1
            return count

        visuals: list[dict[str, Any]] = []
        if not (index == 1 and len(text) < 320):
            for element in elements:
                kind = element.get("type")
                if kind not in ("image", "table"):
                    continue
                x = float(element.get("x") or 0)
                y = float(element.get("y") or 0)
                w = float(element.get("w") or 0)
                h = float(element.get("h") or 0)
                if w < 40 or h < 30:
                    continue
                area = w * h
                area_ratio = area / slide_area
                score = 0.0

                if kind == "table":
                    rows = element.get("rows") or []
                    columns = max((len(row) for row in rows), default=0)
                    non_empty = sum(1 for row in rows for cell in row if str(cell).strip())
                    if len(rows) < 2 or columns < 2 or non_empty < 4:
                        continue
                    if not (0.03 <= area_ratio <= 0.7):
                        continue
                    score = 0.92
                else:
                    aspect = w / h if h else 0
                    if not (0.05 <= area_ratio <= 0.50):
                        continue
                    if aspect < 0.4 or aspect > 3.2:
                        continue
                    if w < DESIGN_WIDTH * 0.16 or h < slide_height_px * 0.08:
                        continue
                    if y < slide_height_px * 0.075 or y + h > slide_height_px * 0.94:
                        continue
                    edges = sum(
                        (
                            x <= DESIGN_WIDTH * 0.02,
                            y <= slide_height_px * 0.02,
                            x + w >= DESIGN_WIDTH * 0.98,
                            y + h >= slide_height_px * 0.98,
                        )
                    )
                    if edges >= 3:
                        continue
                    nearby = nearby_text_count(x, y, w, h)
                    if nearby == 0:
                        continue
                    area_score = 1.0 if 0.07 <= area_ratio <= 0.42 else 0.6
                    aspect_score = 1.0 if 0.5 <= aspect <= 2.8 else 0.55
                    position_score = (
                        1.0
                        if y >= slide_height_px * 0.10 and y + h <= slide_height_px * 0.92
                        else 0.55
                    )
                    size_score = (
                        1.0
                        if w >= DESIGN_WIDTH * 0.22 and h >= slide_height_px * 0.12
                        else 0.5
                    )
                    score = (
                        0.20 * area_score
                        + 0.12 * aspect_score
                        + 0.14 * position_score
                        + 0.12 * size_score
                        + (0.32 if nearby >= 1 else 0.0)
                        + (0.10 if nearby >= 2 else 0.0)
                    )
                    if score < 0.72:
                        continue

                visual = {
                    "id": f"p{index}_v{len(visuals) + 1}",
                    "kind": "image" if kind == "image" else "table",
                    "page": index,
                    "x": round(x, 1),
                    "y": round(y, 1),
                    "w": round(w, 1),
                    "h": round(h, 1),
                    "area": round(area, 1),
                    "confidence": round(score, 2),
                }
                if element.get("url"):
                    visual["url"] = element["url"]
                if element.get("_media_hash"):
                    visual["_media_hash"] = element["_media_hash"]
                visuals.append(visual)
        visuals.sort(key=lambda item: (item.get("confidence", 0), item.get("area", 0)), reverse=True)
        visuals = visuals[:2]

        slides.append(
            {
                "index": index,
                "title": title,
                "text": text,
                "notes": notes,
                "visuals": visuals,
                "preview": {
                    "type": "slides",
                    "width": round(DESIGN_WIDTH, 1),
                    "height": round(slide_height_px, 1),
                    "background": bg or "#ffffff",
                    "elements": elements,
                },
            }
        )

    repeated_threshold = max(2, math.ceil(len(slides) * 0.2))
    for slide in slides:
        kept: list[dict[str, Any]] = []
        for visual in slide.get("visuals", []):
            digest = visual.pop("_media_hash", None)
            if digest and image_hash_counts.get(digest, 0) >= repeated_threshold:
                continue
            kept.append(visual)
        slide["visuals"] = kept
        for element in slide.get("preview", {}).get("elements", []):
            element.pop("_media_hash", None)

    if not slides:
        raise ParseError("该 PPTX 中没有可解析的幻灯片内容。")

    return {
        "doc_id": doc_id,
        "filename": path.name,
        "kind": "pptx",
        "created_at": _now(),
        "page_count": len(slides),
        "slides": slides,
        "warnings": warnings,
    }


# --------------------------------------------------------------------------- #
# pdf
# --------------------------------------------------------------------------- #
def _pdf_page_visuals(
    page: Any, index: int, doc_id: str, media_dir: Path
) -> list[dict[str, Any]]:
    """Select only high-confidence figures with nearby text; no random page crops."""
    try:
        page_dict = page.get_text("dict")
    except Exception:
        return []
    page_area = max(1.0, page.rect.width * page.rect.height)
    page_width = max(1.0, page.rect.width)
    page_height = max(1.0, page.rect.height)

    text_rects: list[Any] = []
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0 or not block.get("bbox"):
            continue
        try:
            rect = pymupdf.Rect(block["bbox"])
        except Exception:
            continue
        if (
            rect.width >= 20
            and rect.height >= 10
            and rect.y0 >= page_height * 0.07
            and rect.y1 <= page_height * 0.94
        ):
            text_rects.append(rect)

    visuals: list[dict[str, Any]] = []
    for block in page_dict.get("blocks", []):
        if block.get("type") != 1:
            continue
        bbox = block.get("bbox")
        if not bbox or len(bbox) != 4:
            continue
        try:
            rect = pymupdf.Rect(bbox)
        except Exception:
            continue
        if rect.width < 60 or rect.height < 50:
            continue
        area = rect.width * rect.height
        area_ratio = area / page_area
        aspect = rect.width / rect.height if rect.height else 0
        if not (0.05 <= area_ratio <= 0.50):
            continue
        if aspect < 0.4 or aspect > 3.0:
            continue
        if rect.width < page_width * 0.16 or rect.height < page_height * 0.08:
            continue
        if rect.y0 < page_height * 0.07 or rect.y1 > page_height * 0.94:
            continue
        edges = sum(
            (
                rect.x0 <= page_width * 0.02,
                rect.y0 <= page_height * 0.02,
                rect.x1 >= page_width * 0.98,
                rect.y1 >= page_height * 0.98,
            )
        )
        if edges >= 3:
            continue

        margin = 45.0
        expanded = pymupdf.Rect(
            max(0, rect.x0 - margin),
            max(0, rect.y0 - margin),
            min(page_width, rect.x1 + margin),
            min(page_height, rect.y1 + margin),
        )
        nearby = sum(1 for text_rect in text_rects if expanded.intersects(text_rect))
        if nearby == 0:
            continue

        area_score = 1.0 if 0.07 <= area_ratio <= 0.42 else 0.6
        aspect_score = 1.0 if 0.5 <= aspect <= 2.8 else 0.55
        position_score = 1.0 if rect.y0 >= page_height * 0.10 and rect.y1 <= page_height * 0.92 else 0.55
        size_score = 1.0 if rect.width >= page_width * 0.20 and rect.height >= page_height * 0.12 else 0.5
        score = (
            0.24 * area_score
            + 0.12 * aspect_score
            + 0.14 * position_score
            + 0.10 * size_score
            + (0.30 if nearby >= 1 else 0.0)
            + (0.10 if nearby >= 2 else 0.0)
        )
        if score < 0.72:
            continue

        clip = pymupdf.Rect(
            max(0, rect.x0 - 4),
            max(0, rect.y0 - 4),
            min(page_width, rect.x1 + 4),
            min(page_height, rect.y1 + 4),
        )
        try:
            pix = page.get_pixmap(clip=clip, dpi=200, alpha=False)
        except Exception:
            continue
        name = f"p{index}_visual{len(visuals) + 1}.png"
        pix.save(str(media_dir / name))
        digest = hashlib.sha1(pix.samples).hexdigest()
        visuals.append(
            {
                "id": f"p{index}_v{len(visuals) + 1}",
                "kind": "image",
                "page": index,
                "x": round(rect.x0, 1),
                "y": round(rect.y0, 1),
                "w": round(rect.width, 1),
                "h": round(rect.height, 1),
                "area": round(area, 1),
                "confidence": round(score, 2),
                "url": f"/api/doc/{doc_id}/media/{name}",
                "_hash": digest,
            }
        )

    visuals.sort(key=lambda item: (item.get("confidence", 0), item.get("area", 0)), reverse=True)
    return visuals[:2]


def _parse_pdf(
    path: Path, doc_id: str, pages_dir: Path, media_dir: Path
) -> dict[str, Any]:
    try:
        pdf = pymupdf.open(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法读取 PDF：{exc}") from exc

    if pdf.needs_pass:
        pdf.close()
        raise ParseError("该 PDF 已加密，请先移除密码后重试。")
    if pdf.page_count == 0:
        pdf.close()
        raise ParseError("PDF 中没有页面。")
    if pdf.page_count > 200:
        pdf.close()
        raise ParseError("PDF 页数超过 200 页，请拆分后再上传。")

    slides: list[dict[str, Any]] = []
    warnings: list[str] = []
    empty_pages = 0
    try:
        for index, page in enumerate(pdf, start=1):
            text = _clean_line(page.get_text("text", sort=True) or "")
            text = re.sub(r"\n{3,}", "\n\n", text)
            text = re.sub(r"(?m)^\s*\d+\s*/\s*\d+\s*$", "", text)
            if len(text) > MAX_TEXT_PER_SLIDE:
                text = text[:MAX_TEXT_PER_SLIDE] + "…"
            if not text:
                empty_pages += 1
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            title = lines[0][:80] if lines else f"第 {index} 页"

            pix = page.get_pixmap(dpi=100, alpha=False)
            image_name = f"p{index}.png"
            pix.save(str(pages_dir / image_name))
            visuals = _pdf_page_visuals(page, index, doc_id, media_dir)
            slides.append(
                {
                    "index": index,
                    "title": title,
                    "text": text,
                    "notes": "",
                    "visuals": visuals,
                    "preview": {
                        "type": "image",
                        "width": round(page.rect.width, 1),
                        "height": round(page.rect.height, 1),
                        "url": f"/api/doc/{doc_id}/pages/{image_name}",
                    },
                }
            )
        if empty_pages == len(slides):
            warnings.append(
                "未从 PDF 中提取到文字（可能是扫描件/图片型 PDF）。左侧仍可预览，但无法生成基于文本的笔记。"
            )
        elif empty_pages:
            warnings.append(f"有 {empty_pages} 页未提取到文字，可能是扫描页。")
    finally:
        pdf.close()

    pdf_hash_counts: dict[str, int] = {}
    for slide in slides:
        for visual in slide.get("visuals", []):
            digest = visual.get("_hash")
            if digest:
                pdf_hash_counts[digest] = pdf_hash_counts.get(digest, 0) + 1
    repeated_threshold = max(2, math.ceil(len(slides) * 0.2))
    for slide in slides:
        kept: list[dict[str, Any]] = []
        for visual in slide.get("visuals", []):
            digest = visual.pop("_hash", None)
            if digest and pdf_hash_counts.get(digest, 0) >= repeated_threshold:
                continue
            kept.append(visual)
        slide["visuals"] = kept

    return {
        "doc_id": doc_id,
        "filename": path.name,
        "kind": "pdf",
        "created_at": _now(),
        "page_count": len(slides),
        "slides": slides,
        "warnings": warnings,
    }


# --------------------------------------------------------------------------- #
# public
# --------------------------------------------------------------------------- #
def parse_document(path: Path, doc_id: str, doc_dir: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        doc = _parse_pptx(path, doc_id, doc_dir / "media")
    elif suffix == ".pdf":
        doc = _parse_pdf(path, doc_id, doc_dir / "pages", doc_dir / "media")
    elif suffix == ".ppt":
        raise ParseError("暂不支持旧版 .ppt，请在 PowerPoint 中另存为 .pptx 后上传。")
    else:
        raise ParseError("仅支持 .pptx 和 .pdf 文件。")

    doc["stats"] = {
        "pages": doc["page_count"],
        "chars": sum(
            len(s.get("text", "")) + len(s.get("notes", ""))
            for s in doc["slides"]
        ),
        "images": sum(
            1
            for s in doc["slides"]
            for e in s.get("preview", {}).get("elements", [])
            if e.get("type") == "image"
        ),
    }
    return doc
